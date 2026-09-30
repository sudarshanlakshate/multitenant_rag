from pathlib import Path
import logging

logger = logging.getLogger(__name__)

SUPPORTED_OFFICE_EXTENSIONS = {".docx", ".pptx"}


def parse_office(file_path: str) -> str:
    """Extract text from supported Office documents using python-docx / python-pptx."""
    path = Path(file_path)

    try:
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {file_path}")

        if path.suffix.lower() not in SUPPORTED_OFFICE_EXTENSIONS:
            raise ValueError(
                f"Unsupported Office extension: {path.suffix}. "
                f"Supported: {sorted(SUPPORTED_OFFICE_EXTENSIONS)}"
            )

        if path.suffix.lower() == ".docx":
            full_text = _parse_docx(str(path))
        else:
            full_text = _parse_pptx(str(path))

        full_text = "\n".join(
            line for line in full_text.splitlines() if line.strip()
        )

        if not full_text.strip():
            logger.warning(
                "Office parser returned empty text: filename=%s",
                file_path,
            )

        logger.info(
            "Office document parsed: filename=%s characters=%d",
            file_path,
            len(full_text),
        )
        return full_text.strip()

    except Exception:
        logger.exception("Office parsing failed: filename=%s", file_path)
        raise


def _parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file using python-docx."""
    from docx import Document  # type: ignore[import]

    doc = Document(file_path)
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    return "\n".join(paragraphs)


def _parse_pptx(file_path: str) -> str:
    """Extract text from a PPTX file using python-pptx."""
    from pptx import Presentation  # type: ignore[import]

    prs = Presentation(file_path)
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        lines.append(text)
    return "\n".join(lines)
